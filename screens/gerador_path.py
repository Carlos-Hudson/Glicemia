import sqlite3
import os
import platform
from datetime import datetime
from kivymd.app import MDApp
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

class GeradorPath:
    def __init__(self):
        # Localiza a instância do app Kivy para pegar a user_data_dir
        self.app = MDApp.get_running_app()
        self.is_android = 'ANDROID_ARGUMENT' in os.environ or platform.machine() == 'armv7l'
        
        # Caminho do banco SQLite (Onde seu app guarda os dados)
        self.db_path = os.path.join(self.app.user_data_dir, "glicemia.db")
        
        # Define a pasta visível para o usuário/médico
        self.pasta_destino = self._configurar_pasta()

    def _configurar_pasta(self):
        """Define o caminho de exportação sem violar segurança do sistema"""
        if self.is_android:
            from android.storage import primary_external_storage_path
            # No Android 10+, Downloads é o local mais seguro para visibilidade
            path = os.path.join(primary_external_storage_path(), "Download", "Glicemia")
        else:
            # Desktop (Ubuntu, Windows, Mac)
            path = os.path.join(os.path.expanduser('~'), "Documents", "Relatorios_Glicemia")
        
        if not os.path.exists(path):
            os.makedirs(path)
        return path

    def exportar(self):
        """Executa a coleta do banco e gera o .xlsx formatado"""
        try:
            # 1. Coleta de dados
            if not os.path.exists(self.db_path):
                return "Erro: Banco de dados não encontrado."
                
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            # Ajuste 'registros' para o nome exato da sua tabela
            cursor.execute("SELECT data, hora, valor FROM registros_glicemia")
            dados = cursor.fetchall()
            conn.close()

            if not dados:
                return "Nenhum dado registrado para exportar."

            # 2. Configuração do Workbook
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Glicemia"

            # 3. Título Principal (Células mescladas e Azul)
            ws.merge_cells('A1:C1')
            ws['A1'] = "Controle diário de Glicemia"
            
            cor_azul_titulo = "3F51B5" # Azul do KivyMD
            ws['A1'].font = Font(size=14, bold=True, color="FFFFFF")
            ws['A1'].fill = PatternFill(start_color=cor_azul_titulo, end_color=cor_azul_titulo, fill_type="solid")
            ws['A1'].alignment = Alignment(horizontal="center", vertical="center")

            # 4. Cabeçalho das Colunas (Linha 2)
            colunas = ["Data", "Horário", "Valor da Glicemia"]
            ws.append(colunas)
            
            for cell in ws[2]:
                cell.font = Font(size=12, bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="7986CB", end_color="7986CB", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")

            # 5. Inserção dos dados com formatação (mg/dL)
            for r in dados:
                # r[0]=Data, r[1]=Hora, r[2]=Valor
                linha_formatada = [r[0], r[1], f"{r[2]} mg/dL"]
                ws.append(linha_formatada)

            # 6. Estilização de todas as células de dados (Fonte 12 e Centro)
            for row in ws.iter_rows(min_row=3):
                for cell in row:
                    cell.font = Font(size=12)
                    cell.alignment = Alignment(horizontal="center")

            # 7. Ajuste automático de largura das colunas
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    if cell.value:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                ws.column_dimensions[column].width = max_length + 4

            # 8. Salvamento final
            data_str = datetime.now().strftime('%d_%m_%Y_%H%M')
            nome_arquivo = f"Relatorio_Glicemia_{data_str}.xlsx"
            caminho_final = os.path.join(self.pasta_destino, nome_arquivo)
            
            wb.save(caminho_final)
            return f"Exportado com sucesso: {nome_arquivo}"

        except Exception as e:
            return f"Ocorreu um erro: {str(e)}"